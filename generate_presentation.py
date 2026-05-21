import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def add_slide(prs, title, content):
    """Utility function to add a standard slide with title and bullet points."""
    slide_layout = prs.slide_layouts[1] # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = content[0] if content else ""
    
    for point in content[1:]:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0
        
    return slide

def generate_pptx():
    prs = Presentation()
    
    # SLIDE 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "HR Employee Attrition Analytics"
    subtitle.text = "Identifying Workforce Risks & Improving Retention\nPrepared by Data Analytics Team"
    
    # SLIDE 2: Business Problem
    add_slide(prs, "The Business Problem", [
        "High employee turnover is costly and disrupts business operations.",
        "HR leadership lacks visibility into exactly WHY employees are leaving.",
        "We need to transition from reactive exit interviews to proactive retention strategies.",
        "Objective: Build an interactive analytics solution to identify high-risk segments and drive data-backed HR decisions."
    ])
    
    # SLIDE 3: Dataset Overview
    add_slide(prs, "Dataset Overview", [
        "Dataset contains 1,470 employee records.",
        "35 distinct features including demographics, compensation, and performance.",
        "Target Variable: Attrition (Yes/No).",
        "Key Indicators Evaluated: Overtime, Monthly Income, Job Satisfaction, Distance from Home."
    ])
    
    # SLIDE 4: ETL Pipeline & Engineering
    add_slide(prs, "ETL Pipeline & Feature Engineering", [
        "Built a robust Python/Pandas ETL pipeline.",
        "Cleaned and standardized the raw CSV data.",
        "Engineered Advanced Business Features:",
        "  - Age Groups & Income Bands",
        "  - Attrition Risk Score (Weighted heuristic)",
        "  - Tenure Categories & Satisfaction Aggregations",
        "Exported clean data into a modular Star Schema."
    ])
    
    # SLIDE 5: Star Schema Data Model
    add_slide(prs, "Star Schema Data Model", [
        "Designed an enterprise-grade Star Schema optimized for BI tools.",
        "Fact Table: FactEmployeeAttrition (Contains measurable metrics and foreign keys)",
        "Dimension Tables:",
        "  - DimEmployee (Demographics)",
        "  - DimDepartment (Roles & Levels)",
        "  - DimCompensation (Salary & Rates)",
        "  - DimSatisfaction (Survey Results)"
    ])
    
    # SLIDE 6: Advanced SQL Analytics
    add_slide(prs, "SQL & DAX Analytics", [
        "SQL Implementations:",
        "  - Utilized CTEs and Window Functions for Risk Ranking.",
        "  - Aggregated attrition drivers by department and income bands.",
        "DAX Implementations (Power BI):",
        "  - CALCULATE & FILTER for context-aware KPIs.",
        "  - Dynamic measure switching for executive dashboards.",
        "  - Time intelligence and ratio calculations."
    ])
    
    # SLIDE 7: Dashboard Highlights
    add_slide(prs, "Dashboard Executive Summary", [
        "Developed a fully interactive Streamlit web application.",
        "Page 1: Executive Summary (High-level KPIs & distributions)",
        "Page 2: Risk Analysis (Correlations between income, satisfaction, and flight risk)",
        "Page 3: Workforce Insights (Retention matrices and raw data drill-downs)",
        "Filters enabled for Department, Gender, and Age Group."
    ])
    
    # SLIDE 8: Key Finding 1 - Overtime
    add_slide(prs, "Finding 1: The Overtime Impact", [
        "Insight: Employees working overtime have a significantly higher attrition rate (3x higher).",
        "Impact: Burnout is directly leading to turnover and increased recruitment costs.",
        "Recommendation: Mandate strict monitoring of overtime hours.",
        "Action: Introduce a 'Burnout Risk Dashboard' for managers."
    ])
    
    # SLIDE 9: Key Finding 2 - Compensation
    add_slide(prs, "Finding 2: Income Band Disparities", [
        "Insight: Attrition is heavily skewed towards the 'Low Income' segment.",
        "Impact: Negative ROI on entry-level hiring; losing talent before optimal productivity.",
        "Recommendation: Conduct an immediate market compensation review for low-income brackets.",
        "Action: Re-evaluate compensation bands and consider staggered micro-promotions."
    ])
    
    # SLIDE 10: Key Finding 3 - Mid-Tenure Slump
    add_slide(prs, "Finding 3: The Mid-Tenure Slump", [
        "Insight: A pronounced spike in attrition occurs around the 3-5 year mark.",
        "Impact: Losing employees right as they become highly experienced domain experts.",
        "Recommendation: Introduce a 'Career Pathing Initiative' at the 2.5-year mark.",
        "Action: Provide lateral movement opportunities and upskilling budgets."
    ])
    
    # SLIDE 11: Future Improvements
    add_slide(prs, "Future Improvements", [
        "Implement predictive Machine Learning models (Random Forest, XGBoost) for Attrition Probability.",
        "Integrate external macro-economic data (inflation, local job market growth).",
        "Transition the ETL pipeline to a cloud-native orchestrated workflow (e.g., Apache Airflow + Snowflake).",
        "Automate distribution of weekly PDF reports to department heads."
    ])
    
    # SLIDE 12: Conclusion
    add_slide(prs, "Conclusion", [
        "Transitioned raw HR data into a proactive, interactive analytics engine.",
        "Identified clear, actionable drivers of attrition (Overtime, Income, Mid-Tenure).",
        "Provided HR Leadership with the tools necessary to improve retention and employee satisfaction.",
        "Thank You. Questions?"
    ])
    
    prs.save('HR_Attrition_Executive_Presentation.pptx')
    print("Presentation generated successfully: HR_Attrition_Executive_Presentation.pptx")

if __name__ == '__main__':
    generate_pptx()
